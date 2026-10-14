from datetime import date
from decimal import Decimal
from io import BytesIO, StringIO

from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse

from accounts.models import Role, RoleAssignment, User
from inspections.tests import Base as InspBase
from norms.models import Indicator
from org.models import OrgUnit
from portfolio.models import PortfolioEntry

from . import data
from .models import IndicatorValue


class HeatmapTests(InspBase):
    def test_heatmap_row_and_scope(self):
        insp, t1, _ = self.make()
        self.fill(self.head, t1, {self.total: "40", self.part: "18"})
        self.client.force_login(self.expert)
        self.client.post(reverse("inspections:transition", args=[insp.pk, "review"]))
        heat = data.heatmap(OrgUnit.objects.all())
        self.assertEqual([r["unit"] for r in heat["rows"]], [self.dep])
        self.assertIn(self.other, heat["unchecked"])
        self.assertEqual(heat["rows"][0]["cells"]["1.1"]["status"], "mild")
        # Заведующий другой кафедры видит панель только своей кафедры
        self.client.force_login(self.other_head)
        r = self.client.get(reverse("dashboard:index"))
        self.assertEqual(r.context["heat"]["rows"], [])
        self.assertIsNone(r.context["strategy"])
        # Преподаватель панель не видит
        self.client.force_login(self.teacher)
        self.assertEqual(self.client.get(reverse("dashboard:index")).status_code, 403)


class StrategyTests(TestCase):
    def setUp(self):
        call_command("load_initial_data", stdout=StringIO())
        uni = OrgUnit.objects.get(kind="university")
        dep = OrgUnit.objects.filter(kind="department").first()
        self.director = User.objects.create_user("d@x.tj", "p", full_name="Д")
        RoleAssignment.objects.create(user=self.director, role=Role.DIRECTOR, org_unit=uni)
        self.rector = User.objects.create_user("r@x.tj", "p", full_name="Р")
        RoleAssignment.objects.create(user=self.rector, role=Role.RECTOR, org_unit=uni)
        self.ts = []
        for k in range(4):
            t = User.objects.create_user(f"t{k}@x.tj", "p", full_name=f"T{k}")
            RoleAssignment.objects.create(user=t, role=Role.TEACHER, org_unit=dep)
            self.ts.append(t)

    def entry(self, user, kind, status="verified", **kw):
        return PortfolioEntry.objects.create(owner=user, kind=kind, title="x", date_start=kw.pop("d", date(2026, 3, 1)), status=status, **kw)

    def test_auto_values_from_verified_portfolio(self):
        self.entry(self.ts[0], "certification")
        self.entry(self.ts[0], "certification")  # тот же человек — считается один раз
        self.entry(self.ts[1], "certification", status="submitted")  # не подтверждено — не считается
        self.entry(self.ts[2], "mobility")
        self.entry(self.ts[0], "publication", indexing="scopus")
        self.entry(self.ts[1], "publication", indexing="rinc")
        self.entry(self.ts[1], "publication", indexing="wos", d=date(2025, 5, 1))  # другой год
        rows = {r["ind"].code: r for r in data.strategy(2026)}
        self.assertEqual(rows["2.2"]["value"], Decimal("25.0"))
        self.assertEqual(rows["2.3"]["value"], Decimal("25.0"))
        self.assertEqual(rows["3.1"]["value"], Decimal(1))
        self.assertEqual(rows["2.2"]["status"], "severe")  # 25 % при цели 2027 ≥ 30
        self.assertEqual(rows["1.1"]["status"], "nodata")
        self.assertEqual(rows["1.4"]["status"], "nodata")

    def test_manual_value_overrides_and_is_logged(self):
        ind = Indicator.objects.get(node__version__framework__code="tsmu-strategy-2030", code="3.1")
        self.client.force_login(self.director)
        self.client.post(reverse("dashboard:strategy_value", args=[ind.pk]) + "?year=2026", {"value": "23", "note": "Scopus"})
        self.assertEqual(IndicatorValue.objects.get(indicator=ind, year=2026).value, Decimal(23))
        row = next(r for r in data.strategy(2026) if r["ind"].code == "3.1")
        self.assertEqual((row["value"], row["status"], row["auto"]), (Decimal(23), "ok", False))
        self.client.force_login(self.rector)
        self.assertEqual(self.client.post(reverse("dashboard:strategy_value", args=[ind.pk]), {"value": "1"}).status_code, 403)

    def test_exports_open(self):
        from docx import Document
        from openpyxl import load_workbook

        self.client.force_login(self.rector)
        x = self.client.get(reverse("dashboard:xlsx") + "?year=2026")
        self.assertEqual(x.status_code, 200)
        wb = load_workbook(BytesIO(x.content))
        self.assertEqual(wb.sheetnames, ["Тепловая карта", "Стратегия 2030", "Планы исправления", "Опросы студентов"])
        self.assertEqual(wb["Стратегия 2030"].max_row, 17)
        d = self.client.get(reverse("dashboard:docx") + "?year=2026")
        text = "\n".join(p.text for p in Document(BytesIO(d.content)).paragraphs)
        self.assertIn("Годовой отчёт Центра стратегического развития за 2026 год", text)
        self.client.force_login(self.ts[0])
        self.assertEqual(self.client.get(reverse("dashboard:docx")).status_code, 403)


@override_settings(DEBUG=True)
class DemoDashboardTests(TestCase):
    def test_demo_dashboard_renders(self):
        import tempfile
        from pathlib import Path

        call_command("load_initial_data", stdout=StringIO())
        with self.settings(BASE_DIR=Path(tempfile.mkdtemp())):
            call_command("setup_demo", stdout=StringIO())
        self.client.force_login(User.objects.get(email="rector@demo.local"))
        r = self.client.get(reverse("dashboard:index"))
        self.assertEqual(len(r.context["heat"]["rows"]), 1)
        self.assertIn("СТ-1", r.context["heat"]["codes"])  # итог опроса студентов попал в карту
        self.assertContains(r, "c-warn")
