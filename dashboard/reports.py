"""Отчёты по кнопке: Excel (таблицы) и Word (годовой отчёт Центра). Данные — из dashboard.data."""

from io import BytesIO

from django.utils import timezone
from django.utils.translation import gettext as _
from django.utils.translation import gettext_lazy

STATUS_TEXT = {"ok": gettext_lazy("в норме"), "mild": gettext_lazy("лёгкое отклонение"),
               "severe": gettext_lazy("выраженное отклонение"), "unknown": gettext_lazy("порог не задан"),
               "nodata": gettext_lazy("нет данных")}


def _num(v):
    return None if v is None else float(v)


def xlsx(ctx, strategy):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    fills = {"ok": "E2F2E8", "mild": "FBF0D9", "severe": "F8E1DF"}
    wb = Workbook()
    ws = wb.active
    ws.title = _("Тепловая карта")
    heat = ctx["heat"]
    ws.append([_("Кафедра")] + heat["titles"] + [_("Отклонений")])
    for row in heat["rows"]:
        ws.append([row["unit"].name] + [_num(c["value"]) if c else None for c in row["line"]] + [row["bad"]])
        for j, c in enumerate(row["line"], start=2):
            if c and c["status"] in fills:
                ws.cell(ws.max_row, j).fill = PatternFill("solid", fgColor=fills[c["status"]])
    ws.append([])
    ws.append([_("Не проверялись: %(n)s кафедр") % {"n": len(heat["unchecked"])}])

    st = wb.create_sheet(_("Стратегия 2030"))
    st.append(["№", _("Показатель"), _("Ед."), _("Цель 2027"), _("Цель 2030"), _("Значение %(year)s") % {"year": ctx["year"]},
              _("Состояние"), _("Источник")])
    for r in strategy:
        i = r["ind"]
        st.append([i.code, i.title, i.unit, _num(i.threshold_medium), _num(i.threshold_high), _num(r["value"]),
                   str(STATUS_TEXT.get(r["status"], "")), r["source"]])

    pl = wb.create_sheet(_("Планы исправления"))
    pl.append(["№", _("Подразделение"), _("Диагноз"), _("План"), _("Статус"), _("Срок"), _("Просрочен"), _("Ответственный")])
    for p in ctx["plan_list"]:
        pl.append([p.ref, str(p.org_unit), str(p.problem), p.title, str(p.get_state_display()), p.due_date,
                   _("да") if p.overdue else "", str(p.owner or "")])

    sv = wb.create_sheet(_("Опросы студентов"))
    sv.append([_("Опрос"), _("Кафедра"), _("Группа"), _("Ответов"), _("Итоги видны")])
    for c in ctx["campaigns"]:
        sv.append([c.title, str(c.org_unit), c.label, c.response_count, _("да") if c.results_visible else _("нет (мало ответов)")])

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(60, max(10, *(len(str(c.value or "")) for c in col)) + 2)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def docx(ctx, strategy):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    year = ctx["year"]
    doc.add_heading(_("Годовой отчёт Центра стратегического развития за %(year)s год") % {"year": year}, level=0)
    doc.add_paragraph(_("ТГМУ им. Абуали ибни Сино. Сформировано программой «Канон» %(date)s. Цифры — по данным, "
                        "внесённым в программу; текстовый анализ дописывается Центром.")
                      % {"date": f"{timezone.localdate():%d.%m.%Y}"})

    doc.add_heading(_("1. Выполнение Стратегии развития до 2030 года"), level=1)
    t = doc.add_table(rows=1, cols=6)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ["№", _("Показатель"), _("Цель 2027"), _("Цель 2030"), f"{year}", _("Состояние")]):
        c.text = h
    for r in strategy:
        i = r["ind"]
        f = lambda v: "—" if v is None else f"{v.normalize():f}"
        cells = t.add_row().cells
        for c, v in zip(cells, [i.code, i.title, f(i.threshold_medium), f(i.threshold_high), f(r["value"]), str(STATUS_TEXT.get(r["status"], ""))]):
            c.text = v
    filled = sum(1 for r in strategy if r["value"] is not None)
    doc.add_paragraph(_("Значения есть по %(filled)s из %(total)s показателей. Показатели 2.2, 2.3 и 3.1 считаются "
                        "по подтверждённым портфолио преподавателей, зарегистрированных в программе.")
                      % {"filled": filled, "total": len(strategy)})

    doc.add_heading(_("2. Проверки кафедр"), level=1)
    heat = ctx["heat"]
    doc.add_paragraph(_("За %(year)s год назначено проверок: %(planned)s, закрыто: %(closed)s. "
                        "Кафедр с результатами: %(rows)s, не проверялись: %(unchecked)s.")
                      % {"year": year, "planned": ctx["inspections_year"], "closed": ctx["inspections_closed"],
                         "rows": len(heat["rows"]), "unchecked": len(heat["unchecked"])})
    for row in heat["rows"]:
        bad = [f"{code} — {c['value'].normalize():f} %" for code, c in zip(heat["codes"], row["line"])
               if c and c["status"] in ("mild", "severe") and c["value"] is not None]
        doc.add_paragraph(f"{row['unit'].name}: " + (_("отклонения — %(list)s") % {"list": "; ".join(bad)} if bad else _("отклонений нет")), style="List Bullet")

    doc.add_heading(_("3. Диагнозы и планы исправления"), level=1)
    ps = ctx["plans"]
    doc.add_paragraph(_("Всего планов: %(total)s, открыто: %(open)s, закрыто результативно: %(closed)s, "
                        "просрочено: %(overdue)s.")
                      % {"total": ps["total"], "open": ps["open"], "closed": ps["closed"], "overdue": len(ps["overdue"])})
    for name, n in ctx["problems"]:
        doc.add_paragraph(f"{name}: {n}", style="List Bullet")

    doc.add_heading(_("4. Обратная связь студентов"), level=1)
    camps = list(ctx["campaigns"])
    doc.add_paragraph(_("Опросов: %(camps)s, ответов: %(responses)s. Итоги учитываются только там, где ответов "
                        "не меньше порога анонимности.") % {"camps": len(camps), "responses": ctx["responses"]})

    doc.add_heading(_("5. Выводы и предложения"), level=1)
    doc.add_paragraph(_("________________________________________________ (заполняет Центр)"))
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
