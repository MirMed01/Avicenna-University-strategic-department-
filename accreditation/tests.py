from datetime import date, timedelta
from io import BytesIO, StringIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse

from accounts.models import Role, RoleAssignment, User
from norms.models import NormNode
from org.models import OrgUnit

from .models import Assessment, Cycle, Rating
from .readiness import forecast


class FakeA:
    def __init__(self, rating):
        self.rating = rating


def fc(strong=0, sat=0, imp=0, unsat=0, none=0):
    items = [FakeA("strong")] * strong + [FakeA("satisfactory")] * sat + [FakeA("improve")] * imp + \
            [FakeA("unsatisfactory")] * unsat + [FakeA("")] * none
    return forecast(items)


class ForecastTests(TestCase):
    """Правила IAAR (4-е изд., 2025) для программной аккредитации."""

    def test_rules(self):
        self.assertEqual(fc(strong=20, sat=79, imp=5)["years"], 7)          # 4.8 % улучшения, 19 % сильных
        self.assertEqual(fc(strong=10, sat=84, imp=6)["years"], 5)          # 6 %, сильных 10 % (<15) → 5
        self.assertEqual(fc(strong=1, sat=79, imp=20)["years"], 3)          # 20 %
        self.assertEqual(fc(strong=0, sat=90, imp=10)["years"], 1)          # нет сильных → только 1 год
        self.assertEqual(fc(strong=5, sat=55, imp=40)["years"], 1)          # 40 %
        self.assertEqual(fc(strong=10, sat=30, imp=60)["years"], 0)         # ровно 60 % — отказ
        self.assertEqual(fc(strong=50, sat=50, unsat=1)["years"], 0)        # один «неудовлетворительно» — отказ

    def test_unrated_not_counted(self):
        r = fc(strong=20, sat=80, none=24)
        self.assertEqual((r["rated"], r["total"], r["years"]), (100, 124, 7))
        self.assertIsNone(fc(none=5)["years"])


class SeedTests(TestCase):
    def test_iaar_seed_and_three_cycles(self):
        call_command("load_initial_data", stdout=StringIO())
        self.assertEqual(NormNode.objects.filter(version__framework__code="iaar-bmifo", kind="criterion").count(), 124)
        self.assertEqual(NormNode.objects.filter(version__framework__code="iaar-bmifo", kind="area").count(), 9)
        self.assertEqual(Cycle.objects.count(), 3)
        c = Cycle.objects.get(programme__contains="Педиатрия")
        self.assertEqual((c.assessments.count(), c.deadline, c.org_unit.name), (124, date(2027, 4, 21), "Педиатрический факультет"))
        self.assertEqual(c.report_due, date(2027, 4, 21) - timedelta(weeks=8))
        call_command("load_initial_data", stdout=StringIO())
        self.assertEqual(Cycle.objects.count(), 3)


@override_settings(MEDIA_ROOT="/tmp/kanon-test-media")
class FlowTests(TestCase):
    def setUp(self):
        call_command("load_initial_data", stdout=StringIO())
        uni = OrgUnit.objects.get(kind="university")
        self.c = Cycle.objects.get(programme__contains="Лечебное")

        def user(email, role, unit):
            u = User.objects.create_user(email, "p", full_name=email)
            RoleAssignment.objects.create(user=u, role=role, org_unit=unit)
            return u

        self.director = user("d@x.tj", Role.DIRECTOR, uni)
        self.dean = user("dean@x.tj", Role.DEAN, self.c.org_unit)
        other_fac = OrgUnit.objects.filter(kind="faculty").exclude(pk=self.c.org_unit_id).first()
        self.other_dean = user("o@x.tj", Role.DEAN, other_fac)
        self.teacher = User.objects.create_user("t@x.tj", "p", full_name="Т")
        self.a = self.c.assessments.order_by("criterion__order").first()

    def post(self, user, a, **data):
        self.client.force_login(user)
        base = {"rating": "", "evidence": "", "justification": "", "measures": ""}
        base.update(data)
        return self.client.post(reverse("accreditation:assessment", args=[a.pk]), base)

    def test_director_rates_and_measures_required(self):
        self.post(self.director, self.a, rating="improve", evidence="Приказ")
        self.a.refresh_from_db()
        self.assertEqual(self.a.rating, "")  # без мер не сохраняется
        self.post(self.director, self.a, rating="improve", evidence="Приказ", measures="Обновить миссию с работодателями")
        self.a.refresh_from_db()
        self.assertEqual(self.a.rating, Rating.IMPROVE)

    def test_assign_standard_and_responsible_can_edit_own_only(self):
        self.client.force_login(self.director)
        self.client.post(reverse("accreditation:assign", args=[self.c.pk]), {"standard": "1", "responsible": self.teacher.pk})
        mine = self.c.assessments.filter(responsible=self.teacher)
        self.assertEqual(mine.count(), 13)
        r = self.post(self.teacher, mine.first(), rating="satisfactory", evidence="Миссия на сайте")
        self.assertEqual(r.status_code, 302)
        other = self.c.assessments.exclude(responsible=self.teacher).first()
        self.assertEqual(self.post(self.teacher, other, rating="strong").status_code, 403)

    def test_access_by_faculty(self):
        self.client.force_login(self.dean)
        self.assertEqual(self.client.get(reverse("accreditation:detail", args=[self.c.pk])).status_code, 200)
        self.client.force_login(self.other_dean)
        self.assertEqual(self.client.get(reverse("accreditation:detail", args=[self.c.pk])).status_code, 403)
        self.assertEqual(self.post(self.dean, self.a, rating="strong").status_code, 403)  # декан смотрит, не правит

    def test_evidence_file_and_docx(self):
        from docx import Document

        self.post(self.director, self.a, rating="strong", evidence="Стратегия 2030", justification="Миссия утверждена Учёным советом")
        self.client.post(reverse("accreditation:upload", args=[self.a.pk]),
                         {"title": "Миссия ТГМУ", "file": SimpleUploadedFile("m.pdf", b"%PDF-1.4", content_type="application/pdf")})
        self.assertEqual(self.a.files.count(), 1)
        r = self.client.get(reverse("accreditation:docx", args=[self.c.pk]))
        doc = Document(BytesIO(r.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("Оценка: Сильная", text)
        self.assertIn("Приложение: Миссия ТГМУ", text)
        self.assertEqual(len(doc.tables[0].rows), 125)  # шапка + 124 критерия
        self.client.force_login(self.teacher)
        self.assertEqual(self.client.get(reverse("accreditation:file", args=[self.a.files.first().pk])).status_code, 404)

    def test_detail_shows_forecast(self):
        Assessment.objects.filter(cycle=self.c).update(rating="satisfactory")
        Assessment.objects.filter(pk=self.a.pk).update(rating="strong")
        self.client.force_login(self.director)
        r = self.client.get(reverse("accreditation:detail", args=[self.c.pk]))
        self.assertEqual(r.context["fc"]["years"], 5)
        self.assertContains(r, "аккредитация на 5 лет")
