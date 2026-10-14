"""Заготовка самоотчёта в Word: раздел II (по критериям) и раздел IV (таблица самооценки), по структуре IAAR."""

from io import BytesIO

from django.utils import timezone
from django.utils.translation import gettext as _

from .models import Rating
from .readiness import by_standard, forecast, standard_of


def self_assessment_docx(cycle, items):
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(12)
    doc.add_heading(_("Отчёт по самооценке: %(p)s") % {"p": cycle.programme}, level=0)
    doc.add_paragraph(_("Стандарт: %(s)s. Заготовка сформирована программой «Канон» %(d)s; текст разделов I и III "
                        "(введение, SWOT-анализ) и выполнение рекомендаций прошлой аккредитации дописываются вручную.")
                      % {"s": cycle.version, "d": timezone.localdate().strftime("%d.%m.%Y")})

    doc.add_heading(_("II. Соответствие стандартам"), level=1)
    current = None
    for a in items:
        std = standard_of(a)
        if std != current:
            current = std
            doc.add_heading(f"{std.code}. {std.title}", level=2)
        p = doc.add_paragraph()
        p.add_run(f"{a.criterion.code}. {a.criterion.text or a.criterion.title}").bold = True
        doc.add_paragraph(_("Оценка: %(r)s") % {"r": a.get_rating_display() or _("не оценено")})
        if a.evidence:
            doc.add_paragraph(_("Доказательства: %(t)s") % {"t": a.evidence})
        for f in a.files.all():
            doc.add_paragraph(_("Приложение: %(t)s") % {"t": f.title}, style="List Bullet")
        if a.justification:
            doc.add_paragraph(_("Обоснование: %(t)s") % {"t": a.justification})
        if a.measures:
            doc.add_paragraph(_("Меры по улучшению: %(t)s") % {"t": a.measures})

    doc.add_heading(_("Итоги по стандартам"), level=2)
    for row in by_standard(items):
        c = row["counts"]
        doc.add_paragraph(_("%(code)s. %(title)s: сильных — %(s)s, удовлетворительных — %(sat)s, предполагают улучшение — %(i)s, "
                            "неудовлетворительных — %(u)s, не оценено — %(n)s.") % {
            "code": row["std"].code, "title": row["std"].title, "s": c[Rating.STRONG], "sat": c[Rating.SATISFACTORY],
            "i": c[Rating.IMPROVE], "u": c[Rating.UNSATISFACTORY], "n": c[""]})

    doc.add_heading(_("IV. Заключение комиссии по самооценке"), level=1)
    labels = [r.label for r in Rating]
    t = doc.add_table(rows=1, cols=2 + len(labels))
    t.style = "Table Grid"
    for cell, h in zip(t.rows[0].cells, [_("№"), _("Критерий")] + [str(x) for x in labels]):
        cell.text = h
    for a in items:
        cells = t.add_row().cells
        cells[0].text = a.criterion.code
        cells[1].text = a.criterion.text or a.criterion.title
        for i, r in enumerate(Rating.values):
            cells[2 + i].text = "+" if a.rating == r else ""
    fc = forecast(items)
    if fc["rated"]:
        doc.add_paragraph(_("Оценено %(r)s из %(t)s критериев. Ориентир по правилам IAAR: %(l)s (решение принимает Совет IAAR).")
                          % {"r": fc["rated"], "t": fc["total"], "l": fc["label"]})
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
