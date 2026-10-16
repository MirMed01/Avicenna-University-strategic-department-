from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.translation import gettext as _

from accounts.decorators import manager_required
from accounts.models import Role
from accounts.permissions import VIEW_ALL, can_manage, has_role, unit_in_scope, units_in_scope
from core import config
from core.models import log
from org.models import OrgUnit

from .analysis import AnalysisError, analyse, flags, parse_rows, read_limited
from .forms import UploadForm
from .models import ExamAnalysis

LOCAL = (Role.HEAD, Role.DEAN)


def _units(user):
    if has_role(user, *VIEW_ALL):
        return OrgUnit.objects.filter(is_active=True, kind=OrgUnit.Kind.DEPARTMENT).order_by("name")
    return units_in_scope(user, LOCAL).filter(kind=OrgUnit.Kind.DEPARTMENT, is_active=True).order_by("name")


def _can_view(user, a):
    return has_role(user, *VIEW_ALL) or unit_in_scope(user, a.org_unit, LOCAL)


@login_required
def index(request):
    qs = ExamAnalysis.objects.select_related("org_unit")
    if not has_role(request.user, *VIEW_ALL):
        qs = qs.filter(org_unit__in=units_in_scope(request.user, LOCAL))
    s = config.get_many("exam_")
    rows = [(a, a.data.get("alpha") is not None and a.data["alpha"] < float(a.alpha_min(s)),
             sum(1 for it in a.data.get("items", []) if flags(it, s))) for a in qs]
    return render(request, "exams/index.html", {"rows": rows, "can_upload": _units(request.user).exists()})


@login_required
def upload(request):
    units = _units(request.user)
    if not units.exists():
        raise PermissionDenied
    form = UploadForm(request.POST or None, request.FILES or None, units=units)
    if form.is_valid():
        from openpyxl import load_workbook

        try:
            wb = load_workbook(form.cleaned_data["file"], read_only=True, data_only=True)
            data = analyse(*parse_rows(read_limited(wb.worksheets[0])))
        except AnalysisError as e:
            form.add_error("file", str(e))
        except Exception:  # повреждённый или не тот файл
            form.add_error("file", _("Не удалось прочитать файл. Сохраните его в Excel как .xlsx и попробуйте снова."))
        else:
            a = form.save(commit=False)
            a.data, a.created_by = data, request.user
            a.save()
            log(request.user, "create", a, f"Анализ теста «{a}»: {data['n']} студентов, {data['k']} заданий, α={data['alpha']}")
            return redirect("exams:detail", a.pk)
    return render(request, "form.html", {"form": form, "title": _("Анализ теста"),
                                         "hint": _("Скачайте образец файла на странице «Анализ тестов».")})


@login_required
def detail(request, pk):
    a = get_object_or_404(ExamAnalysis.objects.select_related("org_unit"), pk=pk)
    if not _can_view(request.user, a):
        raise PermissionDenied
    s = config.get_many("exam_")
    items = []
    for it in a.data["items"]:
        it = dict(it)
        it["flags"] = flags(it, s)
        items.append(it)
    alpha_min = a.alpha_min(s)
    count = {f: sum(1 for it in items if f in it["flags"]) for f in ("hard", "easy", "weak", "negative", "no_variance", "distractor")}
    return render(request, "exams/detail.html", {
        "a": a, "d": a.data, "items": items, "s": s, "alpha_min": alpha_min, "count": count,
        "alpha_low": a.data["alpha"] is not None and a.data["alpha"] < float(alpha_min),
        "few": a.data["n"] < s["exam_min_students"], "to_review": sum(1 for it in items if it["flags"]),
        "can_delete": can_manage(request.user)})


@manager_required
def delete(request, pk):
    a = get_object_or_404(ExamAnalysis, pk=pk)
    if request.method == "POST":
        log(request.user, "delete", a, f"Удалён анализ теста «{a}»")
        a.delete()
        return redirect("exams:index")
    raise PermissionDenied


@login_required
def template(request):
    """Образец файла: заголовки, строка ключа и три студента."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Ответы"
    ws.append([_("Студент")] + [str(i) for i in range(1, 6)])
    ws.append([_("Ключ")] + list("ABCDA"))
    ws.append(["1", "A", "B", "C", "D", "B"])
    ws.append(["2", "A", "C", "C", "D", "A"])
    ws.append(["3", "B", "B", "A", "D", "A"])
    buf = BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = 'attachment; filename="test-analysis-template.xlsx"'
    return resp
