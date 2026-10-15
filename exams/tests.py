from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook

from accounts.models import Role, RoleAssignment, User
from core import config
from org.models import OrgUnit

from .analysis import AnalysisError, analyse, flags, parse_rows
from .models import ExamAnalysis

SCORED = [["№", "1", "2", "3"], ["s1", 1, 1, 1], ["s2", 1, 1, 0], ["s3", 1, 0, 0], ["s4", 0, 0, 0]]


class AnalysisTests(TestCase):
    def test_hand_calculated_example(self):
        """Пример посчитан вручную: p = .75/.5/.25, KR-20 = 0,75, r₁ (исправленная) = 0,522."""
        d = analyse(*parse_rows(SCORED))
        self.assertEqual([i["p"] for i in d["items"]], [0.75, 0.5, 0.25])
        self.assertEqual(d["alpha"], 0.75)
        self.assertEqual(d["items"][0]["r"], 0.522)
        self.assertEqual((d["n"], d["k"], d["mean_pct"]), (4, 3, 50.0))

    def test_key_mode_finds_wrong_key_and_dead_distractor(self):
        # Задание 2: ключ «C», а сильные выбирают «B» — отрицательная различающая способность. Вариант «D» не выбрал никто.
        rows = [["№", "1", "2", "3"], ["Ключ", "A", "C", "D"],
                ["1", "A", "B", "D"], ["2", "A", "B", "D"], ["3", "B", "C", "A"], ["4", "B", "C", "B"], ["5", "A", "B", "D"]]
        d = analyse(*parse_rows(rows))
        s = config.get_many("exam_")
        item2 = d["items"][1]
        self.assertIn("negative", flags(item2, s))
        item1 = d["items"][0]
        self.assertIn("distractor", flags(item1, s))
        self.assertIn("D", item1["dead"])
        self.assertTrue(d["has_key"])

    def test_bad_files_explained(self):
        with self.assertRaises(AnalysisError):
            parse_rows([["№", "1"], ["s1", 1]])
        with self.assertRaises(AnalysisError):
            parse_rows([["№", "1", "2"], ["s1", "A", "B"], ["s2", "B", "A"]])  # буквы без строки «Ключ»

    def test_thresholds_from_settings(self):
        from core.models import Setting

        item = {"p": 0.95, "r": 0.15, "options": None}
        self.assertEqual(flags(dict(item), config.get_many("exam_")), ["easy", "weak"])
        Setting.objects.create(key="exam_p_max", value="0.97")
        Setting.objects.create(key="exam_disc_min", value="0.10")
        self.assertEqual(flags(dict(item), config.get_many("exam_")), [])


def xlsx(rows):
    wb = Workbook()
    for r in rows:
        wb.active.append(r)
    buf = BytesIO()
    wb.save(buf)
    return SimpleUploadedFile("t.xlsx", buf.getvalue())


class UploadTests(TestCase):
    def setUp(self):
        uni = OrgUnit.objects.create(kind="university", name="ТГМУ")
        fac = OrgUnit.objects.create(kind="faculty", name="Мед", parent=uni)
        self.dep = OrgUnit.objects.create(kind="department", name="Хирургия", parent=fac)
        self.other = OrgUnit.objects.create(kind="department", name="Терапия", parent=fac)
        self.head = self.user("h@x.tj", Role.HEAD, self.dep)
        self.other_head = self.user("o@x.tj", Role.HEAD, self.other)
        self.teacher = self.user("t@x.tj", Role.TEACHER, self.dep)

    def user(self, email, role, unit):
        u = User.objects.create_user(email, "pass12345!", full_name=email)
        RoleAssignment.objects.create(user=u, role=role, org_unit=unit)
        return u

    def post(self, unit, rows):
        return self.client.post(reverse("exams:upload"), {
            "title": "Хирургия, 4 курс", "org_unit": unit.pk, "kind": "final", "exam_date": "2026-06-10",
            "file": xlsx(rows)})

    def test_head_uploads_and_student_ids_not_stored(self):
        self.client.force_login(self.head)
        rows = [["Студент"] + SCORED[0][1:]] + [["Иванов И."] + r[1:] for r in SCORED[1:]]
        r = self.post(self.dep, rows)
        a = ExamAnalysis.objects.get()
        self.assertRedirects(r, reverse("exams:detail", args=[a.pk]))
        self.assertNotIn("Иванов", str(a.data))
        r = self.client.get(reverse("exams:detail", args=[a.pk]))
        self.assertContains(r, "0,75")  # русская запись числа
        self.assertContains(r, "меньше 30")  # мало студентов
        self.assertContains(r, "Надёжность теста ниже 0,80")

    def test_head_cannot_upload_for_other_department_or_view_it(self):
        self.client.force_login(self.head)
        r = self.post(self.other, SCORED)
        self.assertEqual((r.status_code, ExamAnalysis.objects.count()), (200, 0))
        a = ExamAnalysis.objects.create(title="T", org_unit=self.other, exam_date="2026-06-10", data=analyse(*parse_rows(SCORED)))
        self.assertEqual(self.client.get(reverse("exams:detail", args=[a.pk])).status_code, 403)

    def test_teacher_cannot_upload(self):
        self.client.force_login(self.teacher)
        self.assertEqual(self.client.get(reverse("exams:upload")).status_code, 403)

    def test_bad_file_shows_reason(self):
        self.client.force_login(self.head)
        r = self.post(self.dep, [["№", "1", "2"], ["s1", "A", "B"], ["s2", "B", "A"]])
        self.assertContains(r, "Без строки «Ключ»")

    def test_template_download(self):
        self.client.force_login(self.head)
        r = self.client.get(reverse("exams:template"))
        self.assertEqual(r.status_code, 200)
        from openpyxl import load_workbook

        rows = list(load_workbook(BytesIO(r.content)).active.iter_rows(values_only=True))
        self.assertEqual(analyse(*parse_rows(rows))["n"], 3)
