from io import StringIO

from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase, override_settings

from accounts.models import User


class DemoSetupTests(TestCase):
    @override_settings(DEBUG=True)
    def test_demo_users_created_and_can_log_in(self):
        call_command("load_initial_data", stdout=StringIO())
        out = StringIO()
        with self.settings(BASE_DIR=self._tmp()):
            call_command("setup_demo", stdout=out)
        password = out.getvalue().strip().rsplit("Пароль: ", 1)[1]
        self.assertEqual(User.objects.filter(email__endswith="@demo.local").count(), 5)
        self.assertTrue(self.client.login(email="teacher@demo.local", password=password))
        teacher = User.objects.get(email="teacher@demo.local")
        self.assertEqual(teacher.roles.get().org_unit.kind, "department")

    @override_settings(DEBUG=False)
    def test_demo_refused_outside_debug(self):
        with self.assertRaises(CommandError):
            call_command("setup_demo", stdout=StringIO())

    def _tmp(self):
        import tempfile
        from pathlib import Path

        d = Path(tempfile.mkdtemp())
        self.addCleanup(lambda: __import__("shutil").rmtree(d, ignore_errors=True))
        return d


class CodespacesLoginTests(TestCase):
    """Codespaces пересылает запрос с Host и Origin = localhost:8000 (https). Вход не должен падать с 403."""

    def test_login_through_codespaces_proxy(self):
        import json

        from django.conf import settings

        env = json.loads((settings.BASE_DIR / ".devcontainer" / "devcontainer.json").read_text(encoding="utf-8"))["containerEnv"]
        origins = env["KANON_CSRF_ORIGINS"].split(",")
        User.objects.create_user("d@demo.local", "Demo-pass-2026", full_name="Д", is_active=True)
        with self.settings(CSRF_TRUSTED_ORIGINS=origins, ALLOWED_HOSTS=["localhost", ".app.github.dev"]):
            c = self.client_class(enforce_csrf_checks=True)
            page = c.get("/login/", HTTP_HOST="localhost:8000")
            token = page.cookies["csrftoken"].value
            r = c.post("/login/", {"username": "d@demo.local", "password": "Demo-pass-2026", "csrfmiddlewaretoken": token},
                       HTTP_HOST="localhost:8000", HTTP_ORIGIN="https://localhost:8000")
        self.assertEqual(r.status_code, 302)


class CodespacesForwardedTests(TestCase):
    """За прокси Codespaces: вход не ломается, а ссылки для студентов ведут на внешний адрес https."""

    def setUp(self):
        import json

        from django.conf import settings

        env = json.loads((settings.BASE_DIR / ".devcontainer" / "devcontainer.json").read_text(encoding="utf-8"))["containerEnv"]
        self.assertEqual(env["KANON_USE_X_FORWARDED_HOST"], "1")
        self.over = self.settings(CSRF_TRUSTED_ORIGINS=env["KANON_CSRF_ORIGINS"].split(","), ALLOWED_HOSTS=["localhost", ".app.github.dev"],
                                  USE_X_FORWARDED_HOST=True, SECURE_PROXY_SSL_HEADER=("HTTP_X_FORWARDED_PROTO", "https"))
        self.over.enable()
        self.addCleanup(self.over.disable)
        self.h = {"HTTP_HOST": "localhost:8000", "HTTP_X_FORWARDED_HOST": "demo-8000.app.github.dev", "HTTP_X_FORWARDED_PROTO": "https"}

    def test_login_and_student_links(self):
        from datetime import date

        from accounts.models import Role, RoleAssignment
        from feedback.models import Campaign
        from instruments.models import Instrument, InstrumentVersion, Item, Section
        from org.models import OrgUnit

        uni = OrgUnit.objects.create(kind="university", name="ТГМУ")
        dep = OrgUnit.objects.create(kind="department", name="К", parent=uni)
        u = User.objects.create_user("d@demo.local", "Demo-pass-2026", full_name="Д", is_active=True)
        RoleAssignment.objects.create(user=u, role=Role.DIRECTOR, org_unit=uni)
        c = self.client_class(enforce_csrf_checks=True)
        token = c.get("/login/", **self.h).cookies["csrftoken"].value
        r = c.post("/login/", {"username": "d@demo.local", "password": "Demo-pass-2026", "csrfmiddlewaretoken": token},
                   HTTP_ORIGIN="https://localhost:8000", **self.h)
        self.assertEqual(r.status_code, 302)
        inst = Instrument.objects.create(code="s", title="А", kind="survey", respondent="student")
        v = InstrumentVersion.objects.create(instrument=inst)
        Item.objects.create(section=Section.objects.create(version=v, title="Р"), code="a", text="?", type="scale")
        v.publish()
        camp = Campaign.objects.create(title="О", version=v, org_unit=dep, closes_on=date(2099, 1, 1))
        camp.add_codes(1)
        page = c.get(f"/feedback/{camp.pk}/codes/", **self.h)
        self.assertContains(page, "https://demo-8000.app.github.dev/s/")


class CsrfFailureTests(TestCase):
    """Дефект, который увидел владелец: устаревшая форма давала английскую техническую страницу 403."""

    def test_stale_form_shows_clear_message(self):
        u = User.objects.create_user("x@demo.local", "Demo-pass-2026", full_name="Х", is_active=True)
        c = self.client_class(enforce_csrf_checks=True)
        c.force_login(u)
        c.get("/")
        r = c.post("/accounts/profile/", {"full_name": "Y", "csrfmiddlewaretoken": "устаревший-ключ" * 2})
        self.assertEqual(r.status_code, 403)
        self.assertContains(r, "Страница устарела", status_code=403)
        self.assertNotContains(r, "CSRF", status_code=403)


@override_settings(DEBUG=True)
class TajikSmokeTests(TestCase):
    """Все основные страницы открываются на таджикском; выгрузки работают; ключевые надписи переведены."""

    @classmethod
    def setUpTestData(cls):
        import tempfile
        from pathlib import Path

        call_command("load_initial_data", stdout=StringIO())
        with override_settings(BASE_DIR=Path(tempfile.mkdtemp())):
            call_command("setup_demo", stdout=StringIO())

    def tg(self, email):
        self.client.force_login(User.objects.get(email=email))
        self.client.post("/i18n/setlang/", {"language": "tg", "next": "/"})

    def test_pages_render_in_tajik(self):
        from accreditation.models import Cycle
        from capa.models import Problem
        from feedback.models import Campaign
        from inspections.models import Inspection
        from instruments.models import Instrument
        from org.models import OrgUnit

        self.tg("director@demo.local")
        insp = Inspection.objects.first()
        urls = ["/", "/dashboard/", "/structure/", f"/structure/{OrgUnit.objects.filter(kind='department').first().pk}/",
                "/norms/", "/instruments/", f"/instruments/{Instrument.objects.first().current_version().pk}/",
                f"/instruments/{Instrument.objects.first().current_version().pk}/try/",
                "/inspections/", f"/inspections/{insp.pk}/", f"/inspections/tasks/{insp.tasks.first().pk}/",
                "/plans/", "/plans/problems/", "/feedback/", f"/feedback/{Campaign.objects.first().pk}/",
                f"/feedback/{Campaign.objects.first().pk}/codes/", "/accreditation/", f"/accreditation/{Cycle.objects.first().pk}/",
                "/portfolio/", "/portfolio/new/", "/accounts/users/", "/journal/", "/inspections/new/", "/feedback/new/"]
        for url in urls:
            r = self.client.get(url)
            self.assertEqual(r.status_code, 200, url)
        home = self.client.get("/").content.decode()
        for ru, tg in (("Структура", "Сохтор"), ("Проверки", "Санҷишҳо")):
            self.assertIn(tg, home)
        self.assertIn('lang="tg"', home)

    def test_exports_in_tajik(self):
        from io import BytesIO

        from openpyxl import load_workbook

        from accreditation.models import Cycle

        self.tg("rector@demo.local")
        x = self.client.get("/dashboard/export.xlsx")
        self.assertEqual(x.status_code, 200)
        names = load_workbook(BytesIO(x.content)).sheetnames
        self.assertTrue(all(len(n) <= 31 for n in names))
        self.assertNotIn("Тепловая карта", names)
        self.assertEqual(self.client.get("/dashboard/report.docx").status_code, 200)
        self.assertEqual(self.client.get(f"/accreditation/{Cycle.objects.first().pk}/report.docx").status_code, 200)

    def test_student_page_in_tajik(self):
        from feedback.models import AccessCode

        code = AccessCode.objects.filter(used=False).first().code
        self.client.post("/i18n/setlang/", {"language": "tg", "next": "/"})
        r = self.client.get(f"/s/{code}/")
        self.assertContains(r, "Машғулиятҳои кафедра")
        self.assertNotContains(r, "Отправить анонимно")


class SeedFitsDatabaseTests(TestCase):
    """Дефект, найденный на PostgreSQL: ссылка IAAR длиннее 200 символов. SQLite длину не проверяет — проверяем сами."""

    def test_all_seed_values_fit_field_lengths(self):
        from django.apps import apps

        call_command("load_initial_data", stdout=StringIO())
        too_long = []
        for model in apps.get_models():
            for f in model._meta.concrete_fields:
                if getattr(f, "max_length", None) and f.get_internal_type() in ("CharField", "URLField", "SlugField", "EmailField", "FileField"):
                    for pk, v in model.objects.values_list("pk", f.attname):
                        if v and len(str(v)) > f.max_length:
                            too_long.append(f"{model.__name__}.{f.name} #{pk}: {len(str(v))} > {f.max_length}")
        self.assertEqual(too_long, [])


class HealthTests(TestCase):
    def test_health_without_login(self):
        r = self.client.get("/health/")
        self.assertEqual((r.status_code, r.content), (200, b"ok"))
